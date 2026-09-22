"""Convert a completed AND/OR Word review into a logic Gold file."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any, Mapping


HEADERS = [
    "Candidate",
    "Target Fault",
    "child_set_complete",
    "logic_gate",
    "evidence_support",
    "overall_decision",
]
ALLOWED_CHILD_SET = {"complete", "incomplete", "unknown"}
ALLOWED_GATES = {"AND", "OR", "unknown", "not_applicable"}
ALLOWED_EVIDENCE = {"supported", "partial", "unsupported"}
ALLOWED_DECISIONS = {"approve", "revise", "reject", "cannot_determine"}


def load_json(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"expected JSON object: {path}")
    return payload


def _clean(value: Any) -> str:
    return str(value or "").strip()


def extract_review_doc(path: Path) -> tuple[list[dict[str, str]], dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - environment-specific
        raise RuntimeError("python-docx is required to read the review document") from exc

    document = Document(path)
    if not document.tables:
        raise ValueError("AND/OR review document must contain a review table")
    table = document.tables[0]
    headers = [_clean(cell.text) for cell in table.rows[0].cells]
    if headers[: len(HEADERS)] != HEADERS or len(headers) != 7:
        raise ValueError(f"unexpected AND/OR review table headers: {headers}")

    rows: list[dict[str, str]] = []
    for row in table.rows[1:]:
        values = [_clean(cell.text) for cell in row.cells]
        if len(values) != 7:
            raise ValueError(f"AND/OR review row has {len(values)} cells instead of 7")
        rows.append(
            {
                "logic_candidate_id": values[0],
                "target_fault": values[1],
                "child_set_complete": values[2],
                "logic_gate": values[3],
                "evidence_support": values[4],
                "overall_decision": values[5],
                "expert_comment": values[6],
            }
        )

    paragraphs = "\n".join(_clean(paragraph.text) for paragraph in document.paragraphs)
    reviewer_match = re.search(r"审核专家\s*[:：]\s*([^\n]+)", paragraphs)
    if reviewer_match is None:
        reviewer_match = re.search(r"专家审核\s*[:：]\s*([^\n]+)", paragraphs)
    date_match = re.search(r"审核日期\s*[:：]\s*(\d{4}-\d{2}-\d{2})", paragraphs)
    summary: dict[str, int] = {}
    for label in ("OR", "AND", "unknown"):
        match = re.search(rf"{re.escape(label)}\s*确认?\s*[:：]\s*(\d+)", paragraphs)
        if match:
            summary[label] = int(match.group(1))
    return rows, {
        "reviewer": _clean(reviewer_match.group(1)) if reviewer_match else None,
        "reviewed_at": date_match.group(1) if date_match else None,
        "review_date_missing": date_match is None,
        "claimed_summary": summary,
    }


def build_gold(
    candidates: Mapping[str, Any],
    reviews: list[Mapping[str, str]],
    metadata: Mapping[str, Any],
    source_review_doc: str,
    source_candidate_bundle: str,
) -> dict[str, Any]:
    candidate_map = {event["logic_candidate_id"]: event for event in candidates["events"]}
    review_map = {review["logic_candidate_id"]: review for review in reviews}
    if set(candidate_map) != set(review_map):
        raise ValueError(
            "review IDs do not match logic candidates; "
            f"missing={sorted(set(candidate_map) - set(review_map))}, "
            f"extra={sorted(set(review_map) - set(candidate_map))}"
        )

    reviewer = _clean(metadata.get("reviewer")) or _clean(
        candidates["dataset_info"].get("expected_reviewer")
    )
    if not reviewer:
        raise ValueError("reviewer is required")

    events: list[dict[str, Any]] = []
    for candidate_id, candidate in candidate_map.items():
        review = review_map[candidate_id]
        if review["target_fault"] != candidate["event_node"]["fault_code"]:
            raise ValueError(f"{candidate_id} target fault does not match candidate bundle")
        for field, allowed in (
            ("child_set_complete", ALLOWED_CHILD_SET),
            ("logic_gate", ALLOWED_GATES),
            ("evidence_support", ALLOWED_EVIDENCE),
            ("overall_decision", ALLOWED_DECISIONS),
        ):
            if review[field] not in allowed:
                raise ValueError(f"{candidate_id} has invalid {field}: {review[field]}")
        if review["overall_decision"] == "approve" and not (
            review["child_set_complete"] == "complete"
            and review["logic_gate"] in {"AND", "OR"}
            and review["evidence_support"] == "supported"
        ):
            raise ValueError(f"{candidate_id} cannot be approved with incomplete gate evidence")

        events.append(
            {
                "logic_candidate_id": candidate_id,
                "event_node": candidate["event_node"],
                "children": candidate["children"],
                "logic_gate": review["logic_gate"],
                "expert_review": {
                    "status": "expert_reviewed",
                    "child_set_complete": review["child_set_complete"],
                    "logic_gate": review["logic_gate"],
                    "evidence_support": review["evidence_support"],
                    "overall_decision": review["overall_decision"],
                    "reviewer": reviewer,
                    "reviewed_at": metadata.get("reviewed_at"),
                    "expert_reason": review["expert_comment"],
                },
            }
        )

    approved = [event for event in events if event["expert_review"]["overall_decision"] == "approve"]
    unknown = [event for event in events if event["logic_gate"] == "unknown"]
    gates_complete = len(unknown) == 0 and all(
        event["expert_review"]["overall_decision"] == "approve" for event in events
    )
    return {
        "dataset_info": {
            "name": "siemens_s210_and_or_logic_gold",
            "version": "v1",
            "status": "expert_validated",
            "gold_source": "named_expert_review",
            "reviewer": reviewer,
            "reviewed_at": metadata.get("reviewed_at"),
            "review_date_missing": bool(metadata.get("review_date_missing")),
            "source_candidate_bundle": source_candidate_bundle,
            "source_review_document": source_review_doc,
            "event_count": len(events),
            "approved_event_count": len(approved),
            "or_count": sum(1 for event in events if event["logic_gate"] == "OR"),
            "and_count": sum(1 for event in events if event["logic_gate"] == "AND"),
            "unknown_gate_count": len(unknown),
            "logic_gate_review_complete": True,
            "logic_gates_complete": gates_complete,
            "causal_relations_complete": False,
            "fta_ready": False,
            "runtime_registry_updated": False,
            "training_eligible": False,
            "claimed_summary": metadata.get("claimed_summary", {}),
            "ingestion_note": (
                "The source review document does not contain a review date."
                if metadata.get("review_date_missing")
                else None
            ),
        },
        "events": events,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--candidates", required=True)
    parser.add_argument("--review-docx", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    candidates_path = Path(args.candidates)
    review_path = Path(args.review_docx)
    candidates = load_json(candidates_path)
    reviews, metadata = extract_review_doc(review_path)
    gold = build_gold(candidates, reviews, metadata, review_path.name, candidates_path.name)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(gold, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "output": str(output),
                "events": len(gold["events"]),
                "approved": gold["dataset_info"]["approved_event_count"],
                "unknown_gates": gold["dataset_info"]["unknown_gate_count"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

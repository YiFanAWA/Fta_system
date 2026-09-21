import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[3]
PUBLIC_SOURCES = ROOT / "evaluation" / "quality_eval" / "public_sources"
sys.path.insert(0, str(PUBLIC_SOURCES))

from import_siemens_s210_expert_docx import build_candidate_dataset  # noqa: E402


class SiemensS210ExpertDocxTests(unittest.TestCase):
    def test_imports_completed_docx_as_candidate_and_flags_known_conflicts(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            sample = {
                "dataset_info": {"review_rules": ["规则1"]},
                "samples": [
                    {
                        "sample_id": "SIEMENS_S210_2019_A01706",
                        "input_text": "A01706 source",
                        "weak_record": {"fault_code": "A01706"},
                        "provenance": {"source_url": "https://example.invalid/source"},
                    }
                ],
            }
            sample_path = root / "sample.json"
            sample_path.write_text(json.dumps(sample, ensure_ascii=False), encoding="utf-8")
            docx_path = root / "review.docx"
            document = Document()
            final = document.add_table(rows=1, cols=2)
            final.rows[0].cells[0].text = "字段"
            final.rows[0].cells[1].text = "专家最终值"
            for key, value in [
                ("故障码", "A01706"),
                ("故障现象", "现象"),
                ("主组件", "SI Motion"),
                ("关联组件", "无"),
                ("候选原因", "原因"),
                ("参数", "p9506"),
                ("gate_type", "unknown"),
            ]:
                cells = final.add_row().cells
                cells[0].text, cells[1].text = key, value
            evidence = document.add_table(rows=1, cols=3)
            for cell, value in zip(evidence.rows[0].cells, ["字段", "原文引用", "页码/位置"]):
                cell.text = value
            for key, value in [
                ("故障码", "A01706"),
                ("故障现象", "现象"),
                ("主组件", "SI Motion"),
                ("关联组件", "The drive is stopped by message F01700."),
                ("候选原因", "原因"),
                ("参数", "p9506"),
            ]:
                cells = evidence.add_row().cells
                cells[0].text, cells[1].text, cells[2].text = key, value, "462"
            document.save(docx_path)

            result = build_candidate_dataset(docx_path, sample_path)
            self.assertEqual(1, result["dataset_info"]["validation_summary"]["records_with_flags"])
            self.assertFalse(result["dataset_info"]["human_expert_reviewed"])
            record = result["records"][0]
            self.assertEqual("审核通过", record["review"]["decision"])
            self.assertEqual("A01706", record["expert_candidate"]["fault_code"])
            self.assertEqual("related_component_evidence_conflict", record["validation"]["flags"][0]["code"])


if __name__ == "__main__":
    unittest.main()

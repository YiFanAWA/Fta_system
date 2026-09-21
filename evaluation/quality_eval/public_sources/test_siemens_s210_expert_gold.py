import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[3]
PUBLIC_SOURCES = ROOT / "evaluation" / "quality_eval" / "public_sources"
sys.path.insert(0, str(PUBLIC_SOURCES))

from finalize_siemens_s210_expert_gold import build_reviewed_dataset  # noqa: E402


class SiemensS210ExpertGoldTests(unittest.TestCase):
    def test_confirmation_applies_three_reclassifications(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            candidate = {
                "records": [
                    {
                        "sample_id": "SIEMENS_S210_2019_A01706",
                        "weak_record": {"fault_code": "A01706"},
                        "input_text": "A01706\nCause: Motion monitoring functions.\nThe drive is stopped by message F01700.",
                        "expert_candidate": {
                            "fault_code": "A01706",
                            "component": "SI Motion",
                            "related_components": "无",
                            "description": "SI Motion",
                            "causes": "Motion monitoring functions.",
                            "parameters": "p9506",
                            "gate_type": "unknown",
                        },
                        "evidence_candidate": {
                            "related_component": {"quote": "The drive is stopped by message F01700.", "location": "462–463"}
                        },
                        "provenance": {},
                    }
                ]
            }
            candidate_path = root / "candidate.json"
            candidate_path.write_text(json.dumps(candidate, ensure_ascii=False), encoding="utf-8")
            confirmation_path = root / "confirmation.docx"
            document = Document()
            for code, target in [
                ("A01706", "反应证据"),
                ("A01788", "原因证据"),
                ("F30655", "原因证据"),
            ]:
                document.add_paragraph(code)
                document.add_paragraph("关联组件最终值：无")
                document.add_paragraph(f"→ 移到故障上下文/{target}")
            document.save(confirmation_path)
            result = build_reviewed_dataset(
                candidate,
                confirmation_path,
                reviewer_name="刘武",
                review_date="2026-09-19",
            )
            record = result["samples"][0]
            self.assertTrue(result["gold_status"]["independent_expert_reviewed"])
            self.assertFalse(result["gold_status"]["f1_ready"])
            self.assertEqual("刘武", result["dataset_info"]["review"]["reviewer_name"])
            self.assertEqual("2026-09-19", result["dataset_info"]["review"]["review_date"])
            self.assertTrue(result["dataset_info"]["review"]["identity_verified"])
            self.assertEqual([], record["gold_records"][0]["related_components"])
            self.assertEqual("reaction_context", record["expert_review"]["reclassification"]["target_field"])


if __name__ == "__main__":
    unittest.main()

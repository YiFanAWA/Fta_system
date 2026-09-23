"""Merge the sixth named-expert causal review batch into official Gold v7."""

from __future__ import annotations

import argparse
import copy
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any, Mapping

from docx import Document

from merge_siemens_s210_causal_relation_gold_v3 import (
    _bool,
    _clean,
    _evidence,
    _normalise_text,
    _validate_review_alignment,
    load_json,
)


def _normalise_review(payload: Mapping[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    info = payload.get("dataset_info")
    reviews = payload.get("reviews")
    if not isinstance(info, Mapping) or not isinstance(reviews, list) or not reviews:
        raise ValueError("expert review JSON must contain dataset_info and non-empty reviews")
    if info.get("review_batch") != "v6" or info.get("review_status") != "completed_expert_review":
        raise ValueError("input is not a completed v6 expert review")

    rows: list[dict[str, Any]] = []
    for record in reviews:
        expert = record.get("expert_review")
        if not isinstance(expert, Mapping):
            raise ValueError(f"{record.get('candidate_id')}.expert_review must be an object")
        rows.append(
            {
                "candidate_id": record.get("candidate_id"),
                "fault_code": record.get("fault_code"),
                "fault_description": record.get("fault_description"),
                "candidate_cause": record.get("candidate_cause"),
                "evidence_ref": record.get("evidence_ref"),
                "evidence_text": record.get("evidence_text"),
                "causal_status": expert.get("causal_status"),
                "direction": expert.get("direction"),
                "relation_type": expert.get("relation_type"),
                "fta_eligible": expert.get("fta_eligible"),
                "overall_decision": expert.get("overall_decision"),
                "expert_comment": expert.get("expert_reason", ""),
                "final_cause_text": expert.get("final_cause_text"),
                "reviewer": expert.get("reviewer"),
                "reviewed_at": expert.get("reviewed_at"),
                "evidence_char_position": record.get("evidence_char_position"),
                "full_source_text": record.get("full_source_text"),
            }
        )
    return dict(info), rows


def _review_summary(rows: list[Mapping[str, Any]]) -> dict[str, int]:
    status = Counter(_clean(row.get("causal_status")) for row in rows)
    decisions = Counter(_clean(row.get("overall_decision")) for row in rows)
    return {
        "total": len(rows),
        "causal": status["causal"],
        "associated_only": status["associated_only"],
        "unsupported": status["unsupported"],
        "cannot_determine": status["cannot_determine"],
        "fta_eligible_true": sum(1 for row in rows if _bool(row.get("fta_eligible"))),
        "approve": decisions["approve"],
        "revise": decisions["revise"],
        "reject": decisions["reject"],
    }


def _validate_offsets_and_sources(
    candidates: Mapping[str, Any], rows: list[Mapping[str, Any]]
) -> dict[str, int]:
    candidate_map = {row["candidate_id"]: row for row in candidates["candidates"]}
    verified_offsets = 0
    verified_sources = 0
    for row in rows:
        candidate_id = _clean(row.get("candidate_id"))
        candidate = candidate_map[candidate_id]
        evidence = (candidate.get("evidence") or [{}])[0]
        source = (candidate.get("source_context") or {}).get("input_text")
        span = re.fullmatch(r"(\d+)-(\d+)", _clean(row.get("evidence_char_position")))
        if not span:
            raise ValueError(f"{candidate_id}.evidence_char_position must be start-end")
        start, end = int(span.group(1)), int(span.group(2))
        if start != evidence.get("start") or end != evidence.get("end"):
            raise ValueError(f"{candidate_id} evidence character positions differ from candidate bundle")
        if not isinstance(source, str) or not isinstance(row.get("full_source_text"), str):
            raise ValueError(f"{candidate_id} is missing full source text")
        if source != row["full_source_text"]:
            raise ValueError(f"{candidate_id} full source text differs from candidate bundle")
        if source[start:end] != evidence.get("quote"):
            raise ValueError(f"{candidate_id} evidence quote does not match its character span")
        if _normalise_text(row.get("evidence_text")) != _normalise_text(evidence.get("quote")):
            raise ValueError(f"{candidate_id} evidence quote differs from expert review")
        verified_offsets += 1
        verified_sources += 1
    return {"verified_evidence_spans": verified_offsets, "verified_full_source_texts": verified_sources}


def _validate_markdown(review_md: Path, rows: list[Mapping[str, Any]]) -> None:
    text = review_md.read_text(encoding="utf-8")
    headings = re.findall(r"(?m)^### (CR-CAND-V6-\d+)\b", text)
    expected_ids = [_clean(row.get("candidate_id")) for row in rows]
    if headings != expected_ids:
        raise ValueError("Markdown candidate headings/order do not match the review JSON")
    blocks = re.split(r"(?m)^### (CR-CAND-V6-\d+)\b.*$", text)
    block_map = {blocks[index]: blocks[index + 1] for index in range(1, len(blocks) - 1, 2)}
    for row in rows:
        candidate_id = _clean(row.get("candidate_id"))
        block = block_map[candidate_id]
        decision = _clean(row.get("overall_decision"))
        if not re.search(rf"overall_decision\s*[:：]\s*`{re.escape(decision)}`", block):
            raise ValueError(f"Markdown decision differs from review JSON for {candidate_id}")
        if not re.search(rf"reviewed_at\s*[:：]\s*`?{re.escape(_clean(row.get('reviewed_at')))}", block):
            raise ValueError(f"Markdown review date differs from review JSON for {candidate_id}")


def _validate_docx(review_docx: Path, rows: list[Mapping[str, Any]], summary: Mapping[str, int]) -> None:
    document = Document(review_docx)
    headings = []
    for paragraph in document.paragraphs:
        match = re.match(r"\s*\d+\.\s+(CR-CAND-V6-\d+)\b", paragraph.text)
        if match:
            headings.append(match.group(1))
    expected_ids = [_clean(row.get("candidate_id")) for row in rows]
    if headings != expected_ids:
        raise ValueError("Word candidate headings/order do not match the review JSON")
    if len(document.tables) != 1 + 3 * len(rows):
        raise ValueError("Word review does not have the expected summary + 3 tables per candidate")

    summary_cells = {
        _normalise_text(line.cells[0].text): _normalise_text(line.cells[1].text)
        for line in document.tables[0].rows
    }
    expected_summary = {
        "候选数量": str(summary["total"]),
        "causal": str(summary["causal"]),
        "associated_only": str(summary["associated_only"]),
        "approve": str(summary["approve"]),
        "revise": str(summary["revise"]),
        "reject": str(summary["reject"]),
        "fta_eligible=true": str(summary["fta_eligible_true"]),
    }
    for field, value in expected_summary.items():
        if summary_cells.get(field) != value:
            raise ValueError(f"Word summary {field} differs from review JSON")

    for index, row in enumerate(rows):
        candidate_table = document.tables[1 + 3 * index]
        review_table = document.tables[2 + 3 * index]
        source_table = document.tables[3 + 3 * index]
        candidate_values = {
            _normalise_text(line.cells[0].text): _normalise_text(line.cells[1].text)
            for line in candidate_table.rows
        }
        review_values = {
            _normalise_text(line.cells[0].text): _normalise_text(line.cells[1].text)
            for line in review_table.rows
        }
        expert_fields = {
            "causal_status": row.get("causal_status"),
            "direction": row.get("direction"),
            "relation_type": row.get("relation_type"),
            "fta_eligible": str(_bool(row.get("fta_eligible"))).lower(),
            "overall_decision": row.get("overall_decision"),
            "审核意见": row.get("expert_comment"),
            "reviewer": row.get("reviewer"),
            "reviewed_at": row.get("reviewed_at"),
        }
        for field, expected in expert_fields.items():
            if review_values.get(field) != _normalise_text(expected):
                raise ValueError(f"Word {row['candidate_id']}.{field} differs from review JSON")
        source_text = _normalise_text(source_table.cell(0, 0).text)
        # The JSON is authoritative for text content; Word may wrap/paragraph-format it.
        if source_text != _normalise_text(row.get("full_source_text")):
            raise ValueError(f"Word full source text differs from review JSON for {row['candidate_id']}")
        candidate_checks = {
            "候选原因": row.get("candidate_cause"),
            "证据原文": row.get("evidence_text"),
        }
        for field, expected in candidate_checks.items():
            if candidate_values.get(field) != _normalise_text(expected):
                raise ValueError(f"Word {row['candidate_id']}.{field} differs from review JSON")


def _is_approved(row: Mapping[str, Any]) -> bool:
    return (
        row.get("causal_status") == "causal"
        and row.get("direction") == "source_to_target"
        and row.get("relation_type") == "causes"
        and _bool(row.get("fta_eligible"))
        and row.get("overall_decision") == "approve"
    )


def _relation_from_review(
    candidate: Mapping[str, Any], row: Mapping[str, Any]
) -> dict[str, Any]:
    source = candidate["source_node"]
    target = candidate["target_node"]
    candidate_id = _clean(row["candidate_id"])
    suffix = candidate_id.removeprefix("CR-CAND-V6-")
    final_cause = _clean(row.get("final_cause_text"))
    if final_cause and _normalise_text(final_cause) != _normalise_text(source["text"]):
        raise ValueError(
            f"{candidate_id} is approved with a revised cause text; a separately reviewed evidence mapping is required"
        )
    comment = _clean(row.get("expert_comment"))
    return {
        "relation_id": f"S210-CAUSAL-V6-{suffix}",
        "candidate_id": candidate_id,
        "source_node": {
            "node_id": source["node_id"],
            "node_type": source["node_type"],
            "text": final_cause or source["text"],
        },
        "target_node": {
            "node_id": target["node_id"],
            "node_type": target["node_type"],
            "fault_code": target["fault_code"],
            "description": target["description"],
        },
        "relation_type": row["relation_type"],
        "direction": row["direction"],
        "relation_note": comment,
        "evidence": _evidence(candidate),
        "expert_review": {
            "causal_status": row["causal_status"],
            "fta_eligible": True,
            "overall_decision": row["overall_decision"],
            "expert_comment": comment,
            "reviewer": _clean(row.get("reviewer")),
            "reviewed_at": _clean(row.get("reviewed_at")),
        },
    }


def _excluded_from_review(candidate: Mapping[str, Any], row: Mapping[str, Any]) -> dict[str, Any]:
    evidence = (candidate.get("evidence") or [{}])[0]
    item = {
        "candidate_id": _clean(row.get("candidate_id")),
        "fault_code": candidate["target_node"]["fault_code"],
        "candidate_cause": candidate["source_node"]["text"],
        "causal_status": row["causal_status"],
        "direction": row["direction"],
        "relation_type": row["relation_type"],
        "fta_eligible": _bool(row["fta_eligible"]),
        "decision": row["overall_decision"],
        "expert_comment": _clean(row.get("expert_comment")),
        "reviewer": _clean(row.get("reviewer")),
        "reviewed_at": _clean(row.get("reviewed_at")),
        "evidence": {
            "citation_id": evidence.get("evidence_id"),
            "quote": evidence.get("quote"),
            "start": evidence.get("start"),
            "end": evidence.get("end"),
        },
    }
    if _clean(row.get("final_cause_text")):
        item["final_cause_text_pending_review"] = _clean(row["final_cause_text"])
    return item


def merge_gold(
    previous: Mapping[str, Any],
    candidates: Mapping[str, Any],
    review_payload: Mapping[str, Any],
    source_review_documents: list[str],
    review_md: Path,
    review_docx: Path,
) -> tuple[dict[str, Any], dict[str, Any]]:
    review_info, rows = _normalise_review(review_payload)
    summary = _review_summary(rows)
    if review_info.get("candidate_count") != len(rows):
        raise ValueError("review candidate_count differs from the actual review rows")
    if review_info.get("summary") != {key: value for key, value in summary.items() if key != "total"}:
        raise ValueError("review JSON summary differs from the recomputed per-record decisions")
    reviewer = _clean(review_info.get("reviewer"))
    reviewed_at = _clean(review_info.get("reviewed_at"))
    if not reviewer or not reviewed_at:
        raise ValueError("review JSON must identify the reviewer and review date")
    if any(_clean(row.get("reviewer")) != reviewer for row in rows):
        raise ValueError("per-record reviewer differs from review dataset_info")
    if any(_clean(row.get("reviewed_at")) != reviewed_at for row in rows):
        raise ValueError("per-record review date differs from review dataset_info")

    _validate_review_alignment(candidates, rows)
    source_checks = _validate_offsets_and_sources(candidates, rows)
    _validate_markdown(review_md, rows)
    _validate_docx(review_docx, rows, summary)

    previous_info = dict(previous.get("dataset_info") or {})
    candidate_info = dict(candidates.get("dataset_info") or {})
    if previous_info.get("version") != "v6":
        raise ValueError("official v7 merge requires official Gold v6 as its base")
    if previous_info.get("gold_source") != "named_expert_review":
        raise ValueError("previous Gold is not sourced from named expert review")
    expected_before = 246
    if previous_info.get("reviewed_candidate_count") != expected_before:
        raise ValueError("Gold v6 reviewed_candidate_count is not the expected 246")
    if candidate_info.get("reviewed_candidate_count") != expected_before:
        raise ValueError("v6 candidate bundle does not start after the expected 246 reviewed candidates")
    if candidate_info.get("candidate_count") != len(rows):
        raise ValueError("v6 candidate bundle does not declare the 50-row batch")
    if candidate_info.get("remaining_unreviewed_candidate_count") != 795:
        raise ValueError("v6 candidate bundle remaining count is not the expected 795")
    source_count = previous_info.get("source_candidate_count")
    if not isinstance(source_count, int) or source_count != 1041:
        raise ValueError("Gold v6 must expose the expected 1041 source candidates")

    prior_candidate_ids = {
        _clean(item.get("candidate_id"))
        for item in [*previous.get("relations", []), *previous.get("excluded_candidates", [])]
    }
    new_ids = {_clean(row.get("candidate_id")) for row in rows}
    if len(new_ids) != len(rows) or prior_candidate_ids.intersection(new_ids):
        raise ValueError("v6 review candidate IDs are duplicated or already present in official Gold v6")

    candidate_map = {row["candidate_id"]: row for row in candidates["candidates"]}
    new_relations = [
        _relation_from_review(candidate_map[row["candidate_id"]], row)
        for row in rows
        if _is_approved(row)
    ]
    new_excluded = [
        _excluded_from_review(candidate_map[row["candidate_id"]], row)
        for row in rows
        if not _is_approved(row)
    ]

    relations = copy.deepcopy(list(previous.get("relations", []))) + new_relations
    excluded = copy.deepcopy(list(previous.get("excluded_candidates", []))) + new_excluded
    relation_ids = [_clean(row.get("relation_id")) for row in relations]
    if len(relation_ids) != len(set(relation_ids)):
        raise ValueError("causal relation IDs overlap after v7 merge")
    counts = Counter(row["target_node"]["fault_code"] for row in relations)
    reviewed_count = previous_info["reviewed_candidate_count"] + len(rows)
    review_dates = dict(previous_info.get("review_dates") or {})
    review_dates["batch_v6"] = reviewed_at
    source_documents = list(previous_info.get("source_review_documents") or [])
    for document in source_review_documents:
        if document not in source_documents:
            source_documents.append(document)
    prior_bundles = list(previous_info.get("source_candidate_bundles") or [])
    batch_bundle = "siemens_s210_causal_relation_candidates_v6_remaining_unreviewed.json"
    if batch_bundle not in prior_bundles:
        prior_bundles.append(batch_bundle)

    gold = {
        "dataset_info": {
            **previous_info,
            "version": "v7",
            "status": "expert_validated",
            "expert_validated": True,
            "expert_validation_scope": "reviewed_candidates_only",
            "reviewer": reviewer,
            "reviewed_at": reviewed_at,
            "review_dates": review_dates,
            "candidate_scope": f"{previous_info.get('candidate_scope', '')}_plus_batch_v6",
            "source_candidate_bundles": prior_bundles,
            "source_review_documents": source_documents,
            "candidate_review_count": reviewed_count,
            "reviewed_candidate_count": reviewed_count,
            "unreviewed_candidate_count": source_count - reviewed_count,
            "approved_causal_relation_count": len(relations),
            "excluded_candidate_count": len(excluded),
            "target_fault_count": len(counts),
            "multi_causal_target_fault_count": sum(1 for count in counts.values() if count >= 2),
            "review_summary_v6": summary,
            "causal_relations_complete": False,
            "logic_gates_complete": False,
            "fta_ready": False,
            "runtime_registry_updated": False,
            "training_eligible": False,
            "merge_note": (
                f"The {len(relations)} relations cover only the named-expert-reviewed candidates "
                "through batch v6; they are not a complete causal Gold for all 281 source records."
            ),
        },
        "relations": relations,
        "excluded_candidates": excluded,
    }
    audit = {
        "status": "merged_named_expert_batch_v6",
        "reviewer": reviewer,
        "reviewed_at": reviewed_at,
        "source_artifact_checks": {
            "json_candidate_count": len(rows),
            "markdown_rows_and_decisions_match_json": True,
            "docx_rows_and_decisions_match_json": True,
            **source_checks,
        },
        "review_summary": summary,
        "promoted_candidate_ids": [_clean(row["candidate_id"]) for row in rows if _is_approved(row)],
        "revise_candidate_ids": [
            _clean(row["candidate_id"]) for row in rows if row["overall_decision"] == "revise"
        ],
        "rejected_candidate_ids": [
            _clean(row["candidate_id"]) for row in rows if row["overall_decision"] == "reject"
        ],
        "official_gold_v7_counts": {
            "relations": len(relations),
            "excluded_or_pending": len(excluded),
            "reviewed_candidates": reviewed_count,
            "unreviewed_candidates": source_count - reviewed_count,
            "source_candidates": source_count,
            "fta_ready": False,
        },
    }
    return gold, audit


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--previous", required=True)
    parser.add_argument("--candidates", required=True)
    parser.add_argument("--review-json", required=True)
    parser.add_argument("--review-md", required=True)
    parser.add_argument("--review-docx", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--audit-output", required=True)
    args = parser.parse_args()

    gold, audit = merge_gold(
        load_json(Path(args.previous)),
        load_json(Path(args.candidates)),
        load_json(Path(args.review_json)),
        [Path(args.review_json).name, Path(args.review_md).name, Path(args.review_docx).name],
        Path(args.review_md),
        Path(args.review_docx),
    )
    for output_arg, payload in ((args.output, gold), (args.audit_output, audit)):
        output = Path(output_arg)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

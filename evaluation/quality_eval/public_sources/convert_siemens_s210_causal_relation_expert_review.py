"""Convert the completed Word expert review into a causal Relation Gold file."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any, Mapping


EXPECTED_HEADERS = [
    "ID",
    "Fault",
    "causal_status",
    "direction",
    "relation_type",
    "fta_eligible",
    "decision",
    "expert_comment",
]


def load_json(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"expected JSON object: {path}")
    return payload


def _clean(value: Any) -> str:
    return str(value or "").strip()


def _bool(value: Any) -> bool:
    return _clean(value).casefold() == "true"


def extract_review_doc(path: Path) -> tuple[list[dict[str, str]], dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - environment-specific
        raise RuntimeError("python-docx is required to read the review document") from exc

    document = Document(path)
    if not document.tables:
        raise ValueError("review document has no review table")
    table = document.tables[0]
    headers = [_clean(cell.text) for cell in table.rows[0].cells]
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"unexpected review table headers: {headers}")

    rows: list[dict[str, str]] = []
    for row in table.rows[1:]:
        values = [_clean(cell.text) for cell in row.cells]
        if len(values) != len(EXPECTED_HEADERS):
            raise ValueError(f"review row has {len(values)} cells instead of 8")
        rows.append(dict(zip(EXPECTED_HEADERS, values)))

    paragraphs = "\n".join(_clean(paragraph.text) for paragraph in document.paragraphs)
    reviewer_match = re.search(r"专家审核\s*[:：]\s*([^\n]+)", paragraphs)
    date_match = re.search(r"审核日期\s*[:：]\s*(\d{4}-\d{2}-\d{2})", paragraphs)
    claimed = {}
    for label in ("causal", "associated_only", "cannot_determine"):
        match = re.search(rf"{re.escape(label)}\s*:\s*(\d+)", paragraphs)
        if match:
            claimed[label] = int(match.group(1))
    metadata = {
        "reviewer": _clean(reviewer_match.group(1)) if reviewer_match else None,
        "reviewed_at": date_match.group(1) if date_match else None,
        "claimed_summary": claimed,
    }
    return rows, metadata


def _review_key(row: Mapping[str, str]) -> str:
    return _clean(row.get("ID"))


def _relation_from_candidate(
    candidate: Mapping[str, Any], review: Mapping[str, str], reviewer: str, reviewed_at: str
) -> dict[str, Any]:
    source = candidate["source_node"]
    target = candidate["target_node"]
    evidence = []
    for item in candidate["evidence"]:
        evidence.append(
            {
                "citation_id": item["evidence_id"],
                "source_id": item["source_id"],
                "field": item["field"],
                "quote": item["quote"],
                "start": item.get("start"),
                "end": item.get("end"),
            }
        )
    return {
        "relation_id": f"S210-CAUSAL-{review['ID'].split('-')[-1]}",
        "candidate_id": review["ID"],
        "source_node": {
            "node_id": source["node_id"],
            "node_type": source["node_type"],
            "text": source["text"],
        },
        "target_node": {
            "node_id": target["node_id"],
            "node_type": target["node_type"],
            "fault_code": target["fault_code"],
            "description": target["description"],
        },
        "relation_type": review["relation_type"],
        "direction": review["direction"],
        "relation_note": review["expert_comment"],
        "evidence": evidence,
        "expert_review": {
            "causal_status": review["causal_status"],
            "fta_eligible": _bool(review["fta_eligible"]),
            "overall_decision": review["decision"],
            "expert_comment": review["expert_comment"],
            "reviewer": reviewer,
            "reviewed_at": reviewed_at,
        },
    }


def build_gold(
    candidates: Mapping[str, Any],
    reviews: list[Mapping[str, str]],
    metadata: Mapping[str, Any],
    source_review_doc: str,
) -> dict[str, Any]:
    candidate_map = {item["candidate_id"]: item for item in candidates["candidates"]}
    review_map = {_review_key(row): row for row in reviews}
    expected_ids = set(candidate_map)
    actual_ids = set(review_map)
    if expected_ids != actual_ids:
        missing = sorted(expected_ids - actual_ids)
        extra = sorted(actual_ids - expected_ids)
        raise ValueError(f"review IDs do not match candidates; missing={missing}, extra={extra}")

    reviewer = _clean(metadata.get("reviewer")) or _clean(candidates["dataset_info"].get("expected_reviewer"))
    reviewed_at = _clean(metadata.get("reviewed_at"))
    if not reviewer or not reviewed_at:
        raise ValueError("reviewer and reviewed_at are required")

    relations = []
    excluded = []
    for review in reviews:
        candidate = candidate_map[review["ID"]]
        is_approved_causal = (
            review["causal_status"] == "causal"
            and review["direction"] == "source_to_target"
            and review["relation_type"] == "causes"
            and _bool(review["fta_eligible"])
            and review["decision"] == "approve"
        )
        if is_approved_causal:
            relations.append(_relation_from_candidate(candidate, review, reviewer, reviewed_at))
        else:
            excluded.append(
                {
                    "candidate_id": review["ID"],
                    "fault_code": candidate["source_record"]["fault_code"],
                    "causal_status": review["causal_status"],
                    "direction": review["direction"],
                    "relation_type": review["relation_type"],
                    "fta_eligible": _bool(review["fta_eligible"]),
                    "decision": review["decision"],
                    "expert_comment": review["expert_comment"],
                    "reviewer": reviewer,
                    "reviewed_at": reviewed_at,
                }
            )

    table_summary = {
        "causal": sum(1 for row in reviews if row["causal_status"] == "causal"),
        "associated_only": sum(1 for row in reviews if row["causal_status"] == "associated_only"),
        "cannot_determine": sum(1 for row in reviews if row["causal_status"] == "cannot_determine"),
        "fta_eligible_true": sum(1 for row in reviews if _bool(row["fta_eligible"])),
    }
    claimed_summary = metadata.get("claimed_summary") or {}
    summary_discrepancy = None
    if claimed_summary and any(claimed_summary.get(key) != table_summary[key] for key in table_summary if key in claimed_summary):
        summary_discrepancy = {
            "claimed_in_document": claimed_summary,
            "reconciled_from_review_table": table_summary,
            "resolution": "Use the per-candidate review table as authoritative; keep the discrepancy for audit.",
        }

    source_info = candidates["dataset_info"]
    return {
        "dataset_info": {
            "name": "siemens_s210_causal_relation_gold",
            "version": "v1",
            "status": "expert_validated",
            "gold_source": "named_expert_review",
            "reviewer": reviewer,
            "reviewed_at": reviewed_at,
            "relation_scope": "causal_only",
            "candidate_scope": "pilot_batch_30",
            "source_candidate_bundle": "siemens_s210_causal_relation_candidates_v1.json",
            "source_review_document": source_review_doc,
            "source_record_count": source_info["source_record_count"],
            "candidate_review_count": len(reviews),
            "approved_causal_relation_count": len(relations),
            "excluded_candidate_count": len(excluded),
            "causal_relations_complete": False,
            "logic_gates_complete": False,
            "fta_ready": False,
            "runtime_registry_updated": False,
            "training_eligible": False,
            "summary_discrepancy": summary_discrepancy,
        },
        "relations": relations,
        "excluded_candidates": excluded,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--candidates", required=True)
    parser.add_argument("--review-docx", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    candidates_path = Path(args.candidates)
    review_path = Path(args.review_docx)
    candidates = load_json(candidates_path)
    reviews, metadata = extract_review_doc(review_path)
    gold = build_gold(candidates, reviews, metadata, review_path.name)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(gold, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "output": str(output),
                "relations": len(gold["relations"]),
                "excluded_candidates": len(gold["excluded_candidates"]),
                "summary_discrepancy": bool(gold["dataset_info"]["summary_discrepancy"]),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

"""Build the focused v6 expert confirmation package for two pending revisions."""

from __future__ import annotations

import argparse
import json
from copy import deepcopy
from pathlib import Path
from typing import Any, Mapping


EXPECTED = {
    "CR-CAND-V5-007": "A01691",
    "CR-CAND-V5-021": "A01782",
}


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def build_confirmation(source: Mapping[str, Any]) -> dict[str, Any]:
    revisions = source.get("revisions")
    if not isinstance(revisions, list):
        raise ValueError("revision package must contain revisions")
    selected = [row for row in revisions if row.get("candidate_id") in EXPECTED]
    if {row.get("candidate_id") for row in selected} != set(EXPECTED):
        raise ValueError("expected exactly A01691 and A01782 revision records")

    output_revisions = []
    for row in selected:
        candidate_id = row["candidate_id"]
        original = deepcopy(row["original_candidate"])
        previous = deepcopy(row["previous_expert_review"])
        suggestion = deepcopy(row["revision_request"])
        suggested_evidence = deepcopy(row.get("suggested_evidence"))
        output_revisions.append(
            {
                "candidate_id": candidate_id,
                "fault_code": EXPECTED[candidate_id],
                "source_record": deepcopy(original["source_record"]),
                "target_node": deepcopy(original["target_node"]),
                "original_candidate": deepcopy(original["source_node"]),
                "original_evidence": deepcopy(original["evidence"]),
                "source_context": deepcopy(original["source_context"]),
                "previous_expert_review": previous,
                "suggested_revision": {
                    "cause_text": suggestion.get("suggested_cause_text", ""),
                    "evidence_action": suggestion.get("suggested_evidence_action", ""),
                    "suggested_evidence": suggested_evidence,
                    "must_be_confirmed_or_replaced": bool(
                        suggestion.get("expert_must_confirm_or_replace")
                    ),
                },
                "expert_review": {
                    "status": "pending_expert_revision",
                    "causal_status": "",
                    "direction": "",
                    "relation_type": "",
                    "fta_eligible": None,
                    "overall_decision": "",
                    "final_cause_text": "",
                    "evidence_reference": "",
                    "evidence_text": "",
                    "evidence_start": None,
                    "evidence_end": None,
                    "expert_reason": "",
                    "reviewer": "刘武",
                    "reviewed_at": "",
                },
            }
        )

    return {
        "dataset_info": {
            "name": "siemens_s210_causal_relation_revision_confirmation",
            "version": "v6",
            "status": "pending_expert_revision",
            "review_type": "causal_relation_revision_confirmation",
            "source_revision_package": "siemens_s210_causal_relation_revision_v5.json",
            "candidate_count": 2,
            "candidate_ids": list(EXPECTED),
            "expected_reviewer": "刘武",
            "review_date": "",
            "expert_validated": False,
            "expert_validation_scope": "pending_revision_records_only",
            "training_eligible": False,
            "causal_relations_complete": False,
            "logic_gates_complete": False,
            "fta_ready": False,
            "not_expert_gold": True,
            "merge_gate": [
                "causal_status=causal",
                "direction=source_to_target",
                "relation_type=causes",
                "fta_eligible=true",
                "overall_decision=approve",
                "final cause text and evidence must be confirmed",
            ],
        },
        "instructions": {
            "purpose": "请刘武逐条确认两条 revise 候选，补齐最终原因文本和准确证据。",
            "do_not_auto_approve": True,
            "approve_only_when": "候选语义、因果方向、关系类型、FTA资格和证据均可接受。",
            "revise_behavior": "若仍需修改，填写最终建议并保持 pending，不进入 Gold。",
            "reject_behavior": "若不成立，选择 reject 并填写原因，保留在排除清单。",
        },
        "revisions": output_revisions,
    }


def _md_code(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2)


def build_markdown(package: Mapping[str, Any]) -> str:
    lines = [
        "# Siemens S210 Causal Relation Revision Confirmation v6",
        "",
        "本确认单只处理 A01691 和 A01782 两条待复核 revise 候选。建议修订不是 Gold 结论，必须由刘武专家逐条确认或替换。确认完成前，两条记录不得进入 Causal Relation Gold v7 或 FTA。",
        "",
        "- 审核专家：刘武",
        "- 当前状态：`pending_expert_revision`",
        "- 审核日期：",
        "- 通过门禁：`approve` + `causal` + `source_to_target` + `causes` + `fta_eligible=true` + 最终原因与证据一致",
        "",
        "## 填写说明",
        "",
        "1. 先核对完整原文和当前证据，不要只看建议文本。",
        "2. `approve` 只有在最终原因、证据定位和因果关系都可接受时才可选择。",
        "3. 选择 `revise` 时，必须填写最终原因文本、证据引用和修改理由；记录继续保持 pending。",
        "4. 选择 `reject` 时，必须填写排除理由；记录不会进入 Gold。",
        "",
    ]
    for index, row in enumerate(package["revisions"], start=1):
        previous = row["previous_expert_review"]
        suggestion = row["suggested_revision"]
        lines.extend(
            [
                f"## {index}. {row['fault_code']} {row['candidate_id']}",
                "",
                f"- 故障描述：{row['source_record']['description']}",
                f"- 当前候选原因：`{row['original_candidate']['text']}`",
                f"- 上一轮结论：`{previous['overall_decision']}`",
                f"- 上一轮专家意见：{previous['review_comment']}",
                f"- 建议原因文本：`{suggestion['cause_text']}`",
                f"- 建议证据处理：{suggestion['evidence_action']}",
                "",
                "### 当前证据",
                "",
                "```json",
                _md_code(row["original_evidence"]),
                "```",
                "",
                "### 完整原文",
                "",
                "```text",
                row["source_context"]["input_text"],
                "```",
                "",
                "### 专家最终填写",
                "",
                "| 字段 | 填写值 |",
                "|---|---|",
                "| causal_status |  |",
                "| direction |  |",
                "| relation_type |  |",
                "| fta_eligible |  |",
                "| overall_decision |  |",
                "| final_cause_text |  |",
                "| evidence_reference |  |",
                "| evidence_text |  |",
                "| evidence_start |  |",
                "| evidence_end |  |",
                "| expert_reason |  |",
                "| reviewer | 刘武 |",
                "| reviewed_at |  |",
                "",
            ]
        )
    lines.extend(
        [
            "## 汇总确认",
            "",
            "- 两条记录是否均已完成复核：",
            "- 可进入 Gold v7 的记录：",
            "- 仍需暂缓的记录：",
            "- 专家补充说明：",
        ]
    )
    return "\n".join(lines) + "\n"


def _set_cell_shading(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_borders(cell: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def _set_cell_text(cell: Any, text: str, bold: bool = False) -> None:
    from docx.oxml.ns import qn

    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = 0
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _add_table(doc: Any, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = True
    _set_cell_text(table.rows[0].cells[0], "字段", True)
    _set_cell_text(table.rows[0].cells[1], "填写值", True)
    for cell in table.rows[0].cells:
        _set_cell_shading(cell, "D9EAF7")
        _set_cell_borders(cell)
    for key, value in rows:
        cells = table.add_row().cells
        _set_cell_text(cells[0], key)
        _set_cell_text(cells[1], value)
        for cell in cells:
            _set_cell_borders(cell)
    doc.add_paragraph().paragraph_format.space_after = 0


def build_docx(package: Mapping[str, Any], output: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Siemens S210 Causal Relation Revision Confirmation v6")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("A01691 and A01782 expert review form").italic = True

    doc.add_heading("审核任务", level=1)
    doc.add_paragraph(
        "本确认单只处理两条 pending revise 候选。请刘武专家逐条核对完整原文、候选原因和证据，填写最终结论。建议修订不代表 Gold 结论；在专家确认前，两条记录不得进入 Gold v7 或 FTA。"
    )
    _add_table(
        doc,
        [
            ("审核专家", "刘武"),
            ("审核日期", ""),
            ("当前状态", "pending_expert_revision"),
            ("批准门禁", "approve + causal + source_to_target + causes + fta_eligible=true"),
        ],
    )
    doc.add_heading("填写规则", level=1)
    for text in [
        "只有最终原因、证据定位、因果方向、关系类型和 FTA 资格均可接受时，才填写 approve。",
        "如果仍需修改，填写 revise、最终原因、证据和理由；该记录继续保持 pending。",
        "如果关系不成立，填写 reject 和排除理由；该记录不会进入 Gold。",
        "证据位置必须对应完整原文中的实际语句，不能只引用相邻但不支持候选语义的文本。",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    for index, row in enumerate(package["revisions"], start=1):
        doc.add_page_break()
        doc.add_heading(f"{index} {row['fault_code']} {row['candidate_id']}", level=1)
        doc.add_paragraph(f"故障描述：{row['source_record']['description']}")
        doc.add_paragraph(f"当前候选原因：{row['original_candidate']['text']}")
        doc.add_paragraph(f"上一轮结论：{row['previous_expert_review']['overall_decision']}")
        doc.add_paragraph(f"上一轮专家意见：{row['previous_expert_review']['review_comment']}")
        doc.add_paragraph(f"建议原因文本：{row['suggested_revision']['cause_text']}")
        doc.add_paragraph(f"建议证据处理：{row['suggested_revision']['evidence_action']}")

        doc.add_heading("当前证据", level=2)
        for evidence in row["original_evidence"]:
            doc.add_paragraph(
                f"引用：{evidence.get('evidence_id', '')}\n"
                f"位置：{evidence.get('start', '')}-{evidence.get('end', '')}\n"
                f"原文：{evidence.get('quote', '')}"
            )
        suggested = row["suggested_revision"].get("suggested_evidence")
        if suggested:
            doc.add_paragraph(
                "建议修正证据（仅供核对）："
                f"位置 {suggested.get('start', '')}-{suggested.get('end', '')}；"
                f"原文：{suggested.get('quote', '')}"
            )

        doc.add_heading("完整原文", level=2)
        source = doc.add_paragraph(row["source_context"]["input_text"])
        source.paragraph_format.left_indent = Inches(0.2)

        doc.add_heading("专家最终填写", level=2)
        _add_table(
            doc,
            [
                ("causal_status", ""),
                ("direction", ""),
                ("relation_type", ""),
                ("fta_eligible", ""),
                ("overall_decision", ""),
                ("final_cause_text", ""),
                ("evidence_reference", ""),
                ("evidence_text", ""),
                ("evidence_start", ""),
                ("evidence_end", ""),
                ("expert_reason", ""),
                ("reviewer", "刘武"),
                ("reviewed_at", ""),
            ],
        )

    doc.add_page_break()
    doc.add_heading("汇总确认", level=1)
    _add_table(
        doc,
        [
            ("两条记录是否均已完成复核", ""),
            ("可进入 Gold v7 的记录", ""),
            ("仍需暂缓的记录", ""),
            ("专家补充说明", ""),
        ],
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--json-output", required=True)
    parser.add_argument("--markdown-output", required=True)
    parser.add_argument("--docx-output", required=True)
    args = parser.parse_args()

    package = build_confirmation(load_json(Path(args.source)))
    for path in (Path(args.json_output), Path(args.markdown_output), Path(args.docx_output)):
        path.parent.mkdir(parents=True, exist_ok=True)
    Path(args.json_output).write_text(
        json.dumps(package, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    Path(args.markdown_output).write_text(build_markdown(package), encoding="utf-8")
    build_docx(package, Path(args.docx_output))
    print(
        json.dumps(
            {
                "json": args.json_output,
                "markdown": args.markdown_output,
                "docx": args.docx_output,
                "candidate_count": len(package["revisions"]),
                "status": package["dataset_info"]["status"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

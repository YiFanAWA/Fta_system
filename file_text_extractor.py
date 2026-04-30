from io import BytesIO
import csv
import json
from typing import Optional


TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".log"}
CSV_EXTENSIONS = {".csv", ".tsv"}
JSON_EXTENSIONS = {".json"}
EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}
DOC_EXTENSIONS = {".docx"}
PDF_EXTENSIONS = {".pdf"}


class FileTextExtractionError(ValueError):
    pass


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FileTextExtractionError("无法解码文件内容，请使用UTF-8/GBK编码文本文件")


def _extract_from_docx(raw: bytes) -> str:
    try:
        from docx import Document  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise FileTextExtractionError("解析 .docx 需要安装 python-docx") from exc

    doc = Document(BytesIO(raw))
    lines = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    return "\n".join(lines)


def _extract_from_pdf(raw: bytes) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise FileTextExtractionError("解析 .pdf 需要安装 pypdf") from exc

    reader = PdfReader(BytesIO(raw))
    texts = []
    for page in reader.pages:
        texts.append(page.extract_text() or "")
    return "\n".join(texts)


def _extract_from_csv(filename: str, raw: bytes) -> str:
    text = _decode_text(raw)
    delimiter = "\t" if filename.lower().endswith(".tsv") else ","
    reader = csv.DictReader(text.splitlines(), delimiter=delimiter)

    rows = []
    for row in reader:
        if not isinstance(row, dict):
            continue
        parts = []
        for key, val in row.items():
            k = (key or "").strip()
            v = (val or "").strip()
            if k and v:
                parts.append(f"{k}={v}")
        if parts:
            rows.append("; ".join(parts))

    return "\n".join(rows) if rows else text


def _extract_from_json(raw: bytes) -> str:
    text = _decode_text(raw)
    try:
        payload = json.loads(text)
    except Exception:
        return text

    lines = []
    if isinstance(payload, list):
        for item in payload:
            if isinstance(item, dict):
                parts = []
                for key, val in item.items():
                    if val is None:
                        continue
                    sval = str(val).strip()
                    if not sval:
                        continue
                    parts.append(f"{key}={sval}")
                if parts:
                    lines.append("; ".join(parts))
            elif item is not None:
                sval = str(item).strip()
                if sval:
                    lines.append(sval)
    elif isinstance(payload, dict):
        for key, val in payload.items():
            if isinstance(val, (dict, list)):
                sval = json.dumps(val, ensure_ascii=False)
            else:
                sval = str(val).strip()
            if sval:
                lines.append(f"{key}={sval}")
    else:
        sval = str(payload).strip()
        if sval:
            lines.append(sval)

    return "\n".join(lines) if lines else text


def _extract_from_xlsx(raw: bytes) -> str:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise FileTextExtractionError("解析 .xlsx 需要安装 openpyxl") from exc

    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration:
            continue

        header_names = [str(h).strip() if h is not None else "" for h in headers]
        for row in rows:
            parts = []
            for idx, cell in enumerate(row):
                val = "" if cell is None else str(cell).strip()
                key = header_names[idx] if idx < len(header_names) else f"col_{idx + 1}"
                if key and val:
                    parts.append(f"{key}={val}")
            if parts:
                lines.append("; ".join(parts))

    return "\n".join(lines)


def extract_text_from_uploaded_file(filename: str, raw: bytes) -> str:
    if not raw:
        raise FileTextExtractionError("上传文件为空")

    suffix = ""
    if filename and "." in filename:
        suffix = filename.lower().rsplit(".", 1)[-1]
        suffix = f".{suffix}"

    if suffix in TEXT_EXTENSIONS or not suffix:
        return _decode_text(raw)

    if suffix in CSV_EXTENSIONS:
        return _extract_from_csv(filename, raw)

    if suffix in JSON_EXTENSIONS:
        return _extract_from_json(raw)

    if suffix in EXCEL_EXTENSIONS:
        return _extract_from_xlsx(raw)

    if suffix in DOC_EXTENSIONS:
        return _extract_from_docx(raw)

    if suffix in PDF_EXTENSIONS:
        return _extract_from_pdf(raw)

    raise FileTextExtractionError(
        "暂不支持该文件类型。支持: txt/md/log/csv/tsv/json/xlsx/docx/pdf"
    )
